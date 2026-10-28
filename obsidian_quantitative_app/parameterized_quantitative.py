"""Google Colab用：複数のNEX DE強度CSVを独立して定量計算し、Excelへ統合する。

複数のCSVと共通の定量計算パラメーターExcelをアップロードする。
各CSVの同日QCだけを用いてドリフト補正・定量計算を行う。
出力Excelには、全測定結果を「定量値」シートへ保存する。
さらにsampleと測定日（時刻を除く）が同じ行をまとめ、各元素の中央値を
「定量値_日付別」シートへ保存する。
資料IDの先頭のアポストロフィだけを除去する。
「定量値」では反復測定をまとめず、「定量値_日付別」で中央値にまとめる。
出力名は入力CSVのファイル名を昇順に結合して自動決定する。
例：20261014.csv、20261019.csv → 20261014_20261019_step1_定量値.xlsx
"""
import importlib.util
import subprocess
import sys
if importlib.util.find_spec("openpyxl") is None:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", "openpyxl"])
import csv
import math
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from statistics import median
from threading import Lock
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
ELEMENTS = ("K", "Ca", "Mn", "Fe", "Zn", "Rb", "Sr", "Y", "Zr", "Nb", "Ag")
QUANT_ELEMENTS = ELEMENTS[:-1]
DIRECT_ELEMENTS = {"K", "Ca", "Mn", "Fe"}
AG_NORMALIZED_ELEMENTS = {"Zn", "Rb", "Sr", "Y", "Zr", "Nb"}
QC_MODE = "obart_prec3_air"
QC_SAMPLE_NAME = "QC-2"
REFERENCE_INTENSITY = {
    "K": 4.27826,
    "Ca": 0.15422,
    "Mn": 0.90407,
    "Fe": 3.15459,
    "Zn": 0.01778,
    "Rb": 0.22959,
    "Sr": 0.07396,
    "Y": 0.09792,
    "Zr": 0.17312,
    "Nb": 0.10719,
    "Ag": 1.53429,
}
CALIBRATION_B = {
    "K": 9312.59,
    "Ca": 30145.8,
    "Mn": 746.088,
    "Fe": 4061.51,
    "Zn": 7048.49,
    "Rb": 1360.26,
    "Sr": 1112.8,
    "Y": 898.725,
    "Zr": 810.974,
    "Nb": 737.066,
}
CALIBRATION_C = {
    "K": -1378.82,
    "Ca": -204.571,
    "Mn": -315.439,
    "Fe": -5513.26,
    "Zn": -44.4817,
    "Rb": -16.0436,
    "Sr": -11.9096,
    "Y": -12.6423,
    "Zr": -18.6844,
    "Nb": -31.1933,
}
# 対象元素: (重なり元素, 重なり補正係数)
OVERLAP_CORRECTION = {
    "Y": ("Rb", -0.113036),
    "Zr": ("Sr", -0.121299),
    "Nb": ("Y", -0.161034),
}
_CONVERSION_LOCK = Lock()

def normalize_parameter_text(value) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip()
    return re.sub(r"[\s_・=()（）]+", "", text)

def normalize_sample_id(value) -> str:
    """資料ID先頭のアポストロフィを除去する（例：'3021B-04 → 3021B-04）。

    資料番号のゼロ埋めや先頭数字は変更しない。集計・中央値計算もしない。
    """
    text = str(value or "").strip()
    return text.lstrip("'‘’＇").strip()



def make_output_name(input_names) -> str:
    """複数CSVの名前を昇順でつなぎ「_step1_定量値.xlsx」を付ける。"""
    stems = []
    for name in input_names:
        if Path(name).suffix.lower() != ".csv":
            raise ValueError(f"CSVファイルではありません: {name}")
        stem = Path(name).stem
        # Colabによる再アップロードの末尾「 (1)」などは名称に含めない。
        stem = re.sub(r"\s*\(\d+\)$", "", stem).strip()
        if not stem:
            raise ValueError(f"CSVのファイル名が空です: {name}")
        stems.append(stem)
    if not stems:
        raise ValueError("CSVが選択されていません。")
    return "_".join(sorted(stems)) + "_step1_定量値.xlsx"


def read_numeric(value, description: str) -> float:
    if value is None or str(value).strip() == "":
        raise ValueError(f"{description}が空欄です。")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{description}を数値として読めません: {value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"{description}は有限の数値にしてください。")
    return number

def find_parameter_value(sheet, labels: tuple[str, ...], required=True):
    targets = {normalize_parameter_text(label) for label in labels}
    for row in sheet.iter_rows():
        for cell in row:
            if normalize_parameter_text(cell.value) in targets:
                return sheet.cell(cell.row, cell.column + 1).value
    if required:
        raise ValueError(f"パラメーターファイルに項目がありません: {labels}")
    return ""

def load_calculation_parameters(parameter_bytes: bytes) -> dict:
    workbook = load_workbook(io.BytesIO(parameter_bytes), data_only=True)
    if "定量計算パラメーター" not in workbook.sheetnames:
        raise ValueError(
            "パラメーターファイルに「定量計算パラメーター」シートがありません。\n"
            f"実際のシート: {workbook.sheetnames}"
        )
    sheet = workbook["定量計算パラメーター"]
    institution = str(
        find_parameter_value(sheet, ("機関名",), required=False) or ""
    ).strip()
    version = str(
        find_parameter_value(
            sheet,
            ("パラメーターバージョン",),
            required=False,
        ) or ""
    ).strip()
    qc_sample_name = str(
        find_parameter_value(
            sheet,
            ("ドリフト補正資料", "QC試料名"),
        )
    ).strip()
    qc_mode = str(
        find_parameter_value(sheet, ("測定モード", "QC測定モード"))
    ).strip()
    if not qc_sample_name:
        raise ValueError("ドリフト補正資料（QC試料名）が空欄です。")
    if not qc_mode:
        raise ValueError("測定モードが空欄です。")
    expected_headers = {
        "元素": "element",
        "QC-2基準強度": "reference",
        "検量線定数B": "b",
        "検量線定数C": "c",
        "補正方法": "method",
        "重なり元素": "overlap_element",
        "重なり補正係数": "overlap_coefficient",
    }
    normalized_headers = {
        normalize_parameter_text(header): key
        for header, key in expected_headers.items()
    }
    header_row = None
    column_map = {}
    for row in sheet.iter_rows():
        row_map = {}
        for cell in row:
            normalized = normalize_parameter_text(cell.value)
            if normalized in normalized_headers:
                row_map[normalized_headers[normalized]] = cell.column
        if {"element", "reference", "b", "c", "method"}.issubset(row_map):
            header_row = row[0].row
            column_map = row_map
            break
    if header_row is None:
        raise ValueError(
            "元素パラメーター表の見出しを確認できません。"
            "必要列: 元素、QC-2基準強度、検量線定数B、"
            "検量線定数C、補正方法、重なり補正係数"
        )
    reference_intensity = {}
    calibration_b = {}
    calibration_c = {}
    direct_elements = set()
    ag_normalized_elements = set()
    overlap_correction = {}
    for row_number in range(header_row + 1, sheet.max_row + 1):
        element_value = sheet.cell(
            row_number,
            column_map["element"],
        ).value
        element = str(element_value or "").strip()
        if not element:
            continue
        if element not in ELEMENTS:
            raise ValueError(f"未対応の元素があります: {element}")
        if element in reference_intensity:
            raise ValueError(f"元素が重複しています: {element}")
        reference = read_numeric(
            sheet.cell(row_number, column_map["reference"]).value,
            f"{element}のQC基準強度",
        )
        if reference <= 0:
            raise ValueError(f"{element}のQC基準強度は0より大きくしてください。")
        reference_intensity[element] = reference
        if element != "Ag":
            calibration_b[element] = read_numeric(
                sheet.cell(row_number, column_map["b"]).value,
                f"{element}の検量線定数B",
            )
            calibration_c[element] = read_numeric(
                sheet.cell(row_number, column_map["c"]).value,
                f"{element}の検量線定数C",
            )
        method = unicodedata.normalize(
            "NFKC",
            str(sheet.cell(row_number, column_map["method"]).value or ""),
        ).replace(" ", "").replace("　", "")
        if "ドリフト補正" not in method:
            raise ValueError(
                f"{element}の補正方法に「ドリフト補正」がありません: {method!r}"
            )
        if element != "Ag":
            if "Ag内標準補正" in method:
                ag_normalized_elements.add(element)
            else:
                direct_elements.add(element)
        coefficient_column = column_map.get("overlap_coefficient")
        coefficient_value = (
            sheet.cell(row_number, coefficient_column).value
            if coefficient_column
            else None
        )
        overlap_column = column_map.get("overlap_element")
        overlap_element = (
            str(sheet.cell(row_number, overlap_column).value or "").strip()
            if overlap_column
            else ""
        )
        if not overlap_element:
            overlap_match = re.search(r"([A-Z][a-z]?)重なり補正", method)
            if overlap_match:
                overlap_element = overlap_match.group(1)
        if coefficient_value not in (None, ""):
            coefficient = read_numeric(
                coefficient_value,
                f"{element}の重なり補正係数",
            )
            if not overlap_element:
                raise ValueError(
                    f"{element}には重なり補正係数がありますが、"
                    "重なり元素を確認できません。"
                )
            if overlap_element not in QUANT_ELEMENTS:
                raise ValueError(
                    f"{element}の重なり元素が未対応です: {overlap_element}"
                )
            overlap_correction[element] = (overlap_element, coefficient)
        elif overlap_element:
            raise ValueError(
                f"{element}には重なり元素がありますが、補正係数が空欄です。"
            )
    missing_elements = [
        element for element in ELEMENTS
        if element not in reference_intensity
    ]
    if missing_elements:
        raise ValueError(
            "パラメーター表に不足している元素があります: "
            + ", ".join(missing_elements)
        )
    if direct_elements | ag_normalized_elements != set(QUANT_ELEMENTS):
        raise ValueError("定量対象元素の補正方法をすべて指定してください。")
    return {
        "institution": institution,
        "version": version,
        "qc_sample_name": qc_sample_name,
        "qc_mode": qc_mode,
        "reference_intensity": reference_intensity,
        "calibration_b": calibration_b,
        "calibration_c": calibration_c,
        "direct_elements": direct_elements,
        "ag_normalized_elements": ag_normalized_elements,
        "overlap_correction": overlap_correction,
    }

def apply_calculation_parameters(parameters: dict) -> None:
    global QC_SAMPLE_NAME, QC_MODE
    global REFERENCE_INTENSITY, CALIBRATION_B, CALIBRATION_C
    global DIRECT_ELEMENTS, AG_NORMALIZED_ELEMENTS, OVERLAP_CORRECTION
    QC_SAMPLE_NAME = parameters["qc_sample_name"]
    QC_MODE = parameters["qc_mode"]
    REFERENCE_INTENSITY = parameters["reference_intensity"]
    CALIBRATION_B = parameters["calibration_b"]
    CALIBRATION_C = parameters["calibration_c"]
    DIRECT_ELEMENTS = parameters["direct_elements"]
    AG_NORMALIZED_ELEMENTS = parameters["ag_normalized_elements"]
    OVERLAP_CORRECTION = parameters["overlap_correction"]

def decode_csv(csv_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "cp932"):
        try:
            return csv_bytes.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("CSVの文字コードを判定できませんでした。")

def parse_datetime(value) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value or "").strip()
    for date_format in (
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
    ):
        try:
            return datetime.strptime(text, date_format)
        except ValueError:
            pass
    raise ValueError(f"測定日時を読めません: {value!r}")

def normalize_qc_name(value) -> str:
    text = str(value or "").strip().upper()
    for character in ("-", "－", "−", "‐", "‑", "–", "—", " ", "　"):
        text = text.replace(character, "")
    return text

def is_qc(measurement: dict) -> bool:
    return (
        normalize_qc_name(measurement["sample"])
        == normalize_qc_name(QC_SAMPLE_NAME)
        and measurement["mode"].strip().lower() == QC_MODE.lower()
    )

def extract_measurements(csv_bytes: bytes) -> list[dict]:
    rows = list(csv.reader(io.StringIO(decode_csv(csv_bytes))))
    measurements = []
    in_intensity_block = False
    for row_number, source_row in enumerate(rows, start=1):
        row = source_row + [""] * max(0, 14 - len(source_row))
        # μが文字化けしても「cps/」で強度ブロックを識別する。
        if sum("cps/" in cell.strip().lower() for cell in row) >= 5:
            in_intensity_block = True
            continue
        if not any(cell.strip() for cell in row):
            in_intensity_block = False
            continue
        if not in_intensity_block:
            continue
        if not all(row[column].strip() for column in range(3)):
            continue
        intensities = {}
        for column, element in enumerate(ELEMENTS, start=3):
            text = row[column].strip()
            if not text:
                raise ValueError(
                    f"CSV {row_number}行目の{element}強度が空欄です。"
                )
            try:
                intensities[element] = float(text)
            except ValueError as exc:
                raise ValueError(
                    f"CSV {row_number}行目の{element}強度を数値として読めません。"
                ) from exc
        if not all(math.isfinite(value) for value in intensities.values()):
            raise ValueError(f"CSV {row_number}行目に非有限の強度があります。")
        measurements.append(
            {
                "csv_row": row_number,
                "sample": normalize_sample_id(row[0]),
                "mode": row[1].strip(),
                "measured_at": parse_datetime(row[2]),
                "intensities": intensities,
            }
        )
    if not measurements:
        raise ValueError("CSV内に強度測定データが見つかりません。")
    return measurements

def calculate_quantitative_values(corrected: dict[str, float]) -> dict[str, float]:
    values = {}
    for element in ("K", "Ca", "Mn", "Fe", "Zn", "Rb", "Sr"):
        values[element] = (
            CALIBRATION_B[element] * corrected[element]
            + CALIBRATION_C[element]
        )
    overlap_element, coefficient = OVERLAP_CORRECTION["Y"]
    values["Y"] = (
        CALIBRATION_B["Y"] * corrected["Y"]
        + CALIBRATION_C["Y"]
        + values[overlap_element] * coefficient
    )
    overlap_element, coefficient = OVERLAP_CORRECTION["Zr"]
    values["Zr"] = (
        CALIBRATION_B["Zr"] * corrected["Zr"]
        + CALIBRATION_C["Zr"]
        + values[overlap_element] * coefficient
    )
    overlap_element, coefficient = OVERLAP_CORRECTION["Nb"]
    values["Nb"] = (
        CALIBRATION_B["Nb"] * corrected["Nb"]
        + CALIBRATION_C["Nb"]
        + values[overlap_element] * coefficient
    )
    return values

def calculate_all(measurements: list[dict]):
    qc_by_date = defaultdict(list)
    for measurement in measurements:
        if is_qc(measurement):
            qc_by_date[measurement["measured_at"].date()].append(measurement)
    for qc_list in qc_by_date.values():
        qc_list.sort(key=lambda item: item["measured_at"])
    if not qc_by_date:
        raise ValueError(
            f"QC試料「{QC_SAMPLE_NAME}」かつ測定モード「{QC_MODE}」の"
            "QC測定が見つかりません。"
        )
    coefficient_rows = []
    coefficients_by_date = defaultdict(list)
    for measurement_date in sorted(qc_by_date):
        for qc_number, qc in enumerate(qc_by_date[measurement_date], start=1):
            coefficients = {}
            for element in ELEMENTS:
                qc_intensity = qc["intensities"][element]
                if qc_intensity == 0:
                    raise ValueError(
                        f"{qc['measured_at']}の{element} QC強度が0です。"
                    )
                coefficients[element] = (
                    REFERENCE_INTENSITY[element] / qc_intensity
                )
            coefficients_by_date[measurement_date].append(
                (qc_number, qc, coefficients)
            )
            coefficient_rows.append(
                {
                    "date": measurement_date,
                    "qc_number": qc_number,
                    "qc": qc,
                    "coefficients": coefficients,
                }
            )
    corrected_rows = []
    missing_qc_dates = set()
    for sample in measurements:
        if is_qc(sample):
            continue
        measurement_date = sample["measured_at"].date()
        daily_coefficients = coefficients_by_date.get(measurement_date, [])
        if not daily_coefficients:
            missing_qc_dates.add(measurement_date)
            continue
        for qc_number, qc, coefficients in daily_coefficients:
            drift_corrected = {
                element: sample["intensities"][element] * coefficients[element]
                for element in ELEMENTS
            }
            corrected_ag = drift_corrected["Ag"]
            if corrected_ag == 0:
                raise ValueError(
                    f"{sample['sample']} / QC{qc_number}: "
                    "ドリフト補正後Ag強度が0です。"
                )
            corrected = {}
            for element in ELEMENTS:
                if element in DIRECT_ELEMENTS or element == "Ag":
                    corrected[element] = drift_corrected[element]
                elif element in AG_NORMALIZED_ELEMENTS:
                    corrected[element] = drift_corrected[element] / corrected_ag
            corrected_rows.append(
                {
                    "sample": sample,
                    "date": measurement_date,
                    "qc_number": qc_number,
                    "qc": qc,
                    "corrected": corrected,
                    "quantitative": calculate_quantitative_values(corrected),
                }
            )
    if not corrected_rows:
        raise ValueError("同日のQCを適用できる試料がありません。")
    return (
        coefficient_rows,
        corrected_rows,
        sorted(missing_qc_dates),
    )

def _append_text_safe(sheet, values):
    """資料名などが '=' で始まってもExcelの式に変換しない。"""
    sheet.append(values)
    for cell, value in zip(sheet[sheet.max_row], values):
        if isinstance(value, str):
            cell.data_type = "s"

def _format_table(sheet, datetime_columns=(), date_columns=(), numeric_columns=(), centered=False):
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="游ゴシック", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="游ゴシック", size=11)
            cell.alignment = Alignment(horizontal="center" if centered else None, vertical="center")
    for column in datetime_columns:
        for cell in sheet[column][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
    for column in date_columns:
        for cell in sheet[column][1:]:
            cell.number_format = "yyyy-mm-dd"
    for column, number_format in numeric_columns:
        for cell in sheet[column][1:]:
            cell.number_format = number_format
    for cells in sheet.columns:
        # 日本語をおおむね2文字幅として見積もる。
        width = max(sum(2 if ord(c) > 255 else 1 for c in str(cell.value or "")) for cell in cells)
        sheet.column_dimensions[cells[0].column_letter].width = min(max(width + 2, 13), 65)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

def create_daily_median_sheet(book, original_sheet):
    """sampleと測定日ごとに、各元素の中央値をまとめたシートを作る。"""
    median_sheet = book.create_sheet("定量値_日付別")
    headers = [cell.value for cell in original_sheet[1]]
    _append_text_safe(median_sheet, headers)

    groups = defaultdict(
        lambda: {
            "modes": set(),
            "values": {element: [] for element in QUANT_ELEMENTS},
        }
    )
    for row in original_sheet.iter_rows(min_row=2, values_only=True):
        sample, mode, measured_at, *quantitative_values = row
        measurement_date = parse_datetime(measured_at).date()
        group = groups[(sample, measurement_date)]
        group["modes"].add(mode)
        for element, value in zip(QUANT_ELEMENTS, quantitative_values):
            group["values"][element].append(value)

    for sample, measurement_date in sorted(groups):
        group = groups[(sample, measurement_date)]
        modes = group["modes"]
        mode_label = next(iter(modes)) if len(modes) == 1 else "複数モード"
        _append_text_safe(
            median_sheet,
            [
                sample,
                mode_label,
                measurement_date,
                *[median(group["values"][element]) for element in QUANT_ELEMENTS],
            ],
        )

    _format_table(
        median_sheet,
        date_columns=("C",),
        centered=True,
        numeric_columns=tuple(
            (get_column_letter(i), "0.000") for i in range(4, 14)
        ),
    )
    return median_sheet

def create_multi_quantitative_excel(csv_files, parameter_content, parameter_name=""):
    """アップロード結果 {CSV名: bytes} と共通Excel bytesから統合Excelを返す。
    CSVごとのQC・資料測定・補正係数を隔離し、計算結果だけを連結する。
    不正ファイルがあれば、そのファイルだけ黙って除外せず全体を停止する。
    """
    if not csv_files:
        raise ValueError("強度データのCSVが選択されていません。")
    if not parameter_content:
        raise ValueError("定量計算パラメーターExcelが空です。")
    invalid_names = [name for name in csv_files if Path(name).suffix.lower() != ".csv"]
    if invalid_names:
        raise ValueError("CSV以外のファイルがあります：" + "、".join(invalid_names))
    parameters = load_calculation_parameters(parameter_content)
    batches = []
    with _CONVERSION_LOCK:
        apply_calculation_parameters(parameters)
        for name, content in csv_files.items():
            try:
                if not content:
                    raise ValueError("ファイルが空です。")
                measurements = extract_measurements(content)
                coefficients, corrected, missing = calculate_all(measurements)
                for row in corrected:
                    if not all(math.isfinite(value) for value in row["quantitative"].values()):
                        raise ValueError("計算結果に非有限の数値があります。")
                samples = [m for m in measurements if not is_qc(m)]
                batches.append(dict(name=name, coefficients=coefficients, corrected=corrected,
                                    missing=missing, measurements=len(samples),
                                    excluded=sum(m["measured_at"].date() in missing for m in samples)))
            except (ValueError, KeyError, ZeroDivisionError, OverflowError) as error:
                raise ValueError(f"「{name}」を処理できませんでした：{error}\n"
                                 "不完全な統合結果を出さないため、全体の出力を停止しました。") from error
    book = Workbook()
    quant = book.active
    quant.title = "定量値"
    _append_text_safe(quant, ["sample", "測定モード", "試料測定日時",
                             *[f"{e}_ppm" for e in QUANT_ELEMENTS]])

    # QC補正係数・処理一覧は計算に利用するが、Excelシートとしては保存しない。
    # QCが欠けた測定日の警告はColabの画面に表示する。
    warning_rows = []
    for batch in batches:
        name = batch["name"]
        for row in batch["corrected"]:
            sample = row["sample"]
            _append_text_safe(quant, [sample["sample"], sample["mode"], sample["measured_at"],
                                      *[row["quantitative"][e] for e in QUANT_ELEMENTS]])
        for missing_date in batch["missing"]:
            warning_rows.append([name, missing_date, "同日のQC測定がないため、この測定日の資料を除外"])

    _format_table(quant, datetime_columns=("C",), centered=True,
                  numeric_columns=tuple((get_column_letter(i), "0.000") for i in range(4, 14)))

    median_sheet = create_daily_median_sheet(book, quant)
    book.active = book.sheetnames.index(median_sheet.title)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return dict(content=output.getvalue(), file_count=len(batches),
                output_row_count=sum(len(b["corrected"]) for b in batches),
                median_row_count=median_sheet.max_row - 1,
                warnings=warning_rows, batches=batches)

def main():
    from google.colab import files
    import hashlib
    print("① 強度データのCSVを選択してください。複数ファイルを同時に選択できます。")
    uploaded_csv = files.upload()
    if not uploaded_csv:
        raise ValueError("CSVが選択されていません。セルを再実行してください。")
    if any(Path(name).suffix.lower() != ".csv" for name in uploaded_csv):
        raise ValueError("最初のアップロードではCSVだけを選択してください。")
    print(f"CSV {len(uploaded_csv)}ファイルを受け付けました。")
    seen = {}
    for name, content in uploaded_csv.items():
        digest = hashlib.sha256(content).hexdigest()
        if digest in seen:
            print(f"注意：「{name}」と「{seen[digest]}」の内容が同一です。両方を計算対象にします。")
        seen[digest] = name
    print("② 共通して使う定量計算パラメーターExcelを1つ選択してください。")
    uploaded_parameter = files.upload()
    if len(uploaded_parameter) != 1:
        raise ValueError("定量計算パラメーターExcelは1つだけ選択してください。")
    parameter_name, parameter_content = next(iter(uploaded_parameter.items()))
    if Path(parameter_name).suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("パラメーターはxlsxまたはxlsm形式のExcelを選択してください。")
    print("③ CSVごとに独立してドリフト補正・定量計算を行います。")
    result = create_multi_quantitative_excel(uploaded_csv, parameter_content, parameter_name)
    for batch in result["batches"]:
        print(f"  {batch['name']}：資料測定 {batch['measurements']}件、"
              f"QC {len(batch['coefficients'])}件、定量値出力 {len(batch['corrected'])}行、"
              f"除外測定 {batch['excluded']}件")
    for name, date_value, message in result["warnings"]:
        print(f"警告：{name} / {date_value.isoformat()} / {message}")
    output_name = make_output_name(uploaded_csv.keys())
    # 再実行時も常に同じ名前で出力する。Colab上の同名ファイルは上書きする。
    target = Path(output_name)
    target.write_bytes(result["content"])
    print(f"④ 完了：{result['file_count']}ファイル・詳細{result['output_row_count']}行・"
          f"中央値{result['median_row_count']}行")
    print("ダウンロード：", target.name)
    files.download(str(target))
if __name__ == "__main__":
    main()
