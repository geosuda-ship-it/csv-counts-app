from io import BytesIO

import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Obsidian EDXRF Ellipse Method",
    layout="wide",
)

st.title("判別楕円法による黒曜石の化学組成グループ候補の評価")

uploaded = st.file_uploader(
    "① 測定データのExcelファイルをアップロードしてください。",
    type=["xlsx"],
)

if uploaded is not None:
    try:
        data = pd.read_excel(
            BytesIO(uploaded.getvalue()),
            sheet_name="定量値",
            dtype={"sample": "string"},
            engine="openpyxl",
        ).dropna(how="all").copy()

        if data.empty:
            raise ValueError("定量値シートに測定データがありません。")

        if "sample" not in data.columns:
            raise ValueError("定量値シートにsample列がありません。")

        # 元の資料IDを残す
        data["sample_元表記"] = data["sample"].copy()

        # アポストロフィを無視して資料IDを統一
        data["sample"] = (
            data["sample"]
            .str.replace("'", "", regex=False)
            .str.strip()
        )

        if data["sample"].isna().any() or data["sample"].eq("").any():
            raise ValueError("資料IDが空白の測定データがあります。")

    except Exception as exc:
        st.error(f"読み込めませんでした：{exc}")
        st.stop()

    sample_counts = (
        data.groupby("sample", sort=False)
        .size()
        .reset_index(name="測定回数")
    )

    st.success(
        f"{len(sample_counts)}資料・{len(data)}測定を読み込みました。"
    )

    st.subheader("資料別の測定回数")
    st.dataframe(sample_counts, hide_index=True, width="stretch")

    with st.expander("測定データを確認する"):
        st.dataframe(data, hide_index=True, width="stretch")
if uploaded is not None:
    import openpyxl

    parameter_file = st.file_uploader(
        "② 判別楕円パラメーターのExcelファイルをアップロードしてください。",
        type=["xlsx"],
    )

    if parameter_file is not None:
        parameter_columns = [
            "Diagram", "Group", "X", "Y",
            "x0", "y0", "a", "b", "theta",
        ]

        parameter_rows = []
        workbook = None

        try:
            workbook = openpyxl.load_workbook(
                BytesIO(parameter_file.getvalue()),
                read_only=True,
                data_only=True,
            )

            if "AS-1" not in workbook.sheetnames:
                raise ValueError("AS-1シートがありません。")

            start = workbook.sheetnames.index("AS-1")

            for sheet_name in workbook.sheetnames[start:]:
                sheet = workbook[sheet_name]

                block = list(sheet.iter_rows(
                    min_row=4,
                    max_row=10,
                    min_col=12,
                    max_col=20,
                    values_only=True,
                ))

                headers = [
                    str(value).strip() if value is not None else ""
                    for value in block[0]
                ]

                if headers != parameter_columns:
                    raise ValueError(
                        f"{sheet_name}シートのL4:T4の見出しが一致しません。"
                    )

                group_data = pd.DataFrame(
                    block[1:],
                    columns=parameter_columns,
                )

                if sorted(group_data["Diagram"].tolist()) != [1, 2, 3, 4, 5, 6]:
                    raise ValueError(
                        f"{sheet_name}シートにD1〜D6が揃っていません。"
                    )

                if not group_data["Group"].eq(sheet_name).all():
                    raise ValueError(
                        f"{sheet_name}シートのGroupとシート名が一致しません。"
                    )

                parameter_rows.extend(
                    group_data.to_dict("records")
                )

        except Exception as exc:
            st.error(f"パラメーターを読み込めませんでした：{exc}")
            st.stop()

        finally:
            if workbook is not None:
                workbook.close()

        ellipse_parameters = pd.DataFrame(
            parameter_rows,
            columns=parameter_columns,
        )

        st.success(
            f"{ellipse_parameters['Group'].nunique()}化学組成グループ・"
            f"{len(ellipse_parameters)}楕円を読み込みました。"
        )

        with st.expander("判別楕円パラメーターを確認する"):
            st.dataframe(
                ellipse_parameters,
                hide_index=True,
                width="stretch",
            )
