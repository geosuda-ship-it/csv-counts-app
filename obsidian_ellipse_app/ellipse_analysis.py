"""Input stage of the ellipse method. Each source row is one measurement."""

from dataclasses import dataclass

from datetime import date, datetime

from io import BytesIO

from math import isfinite

from numbers import Real

from zipfile import BadZipFile



import openpyxl

import pandas as pd





ELEMENTS = ("K", "Ca", "Mn", "Fe", "Zn", "Rb", "Sr", "Y", "Zr", "Nb")

REQUIRED_COLUMNS = ("sample", "測定モード", "試料測定日時") + tuple(f"{e}_ppm" for e in ELEMENTS)

PARAMETER_COLUMNS = ("Diagram", "Group", "X", "Y", "x0", "y0", "a", "b", "theta")

DIAGRAMS = {

    1: ("Rb*", "Sr*"),

    2: ("Rb*", "Zr*"),

    3: ("Sr*", "Zr*"),

    4: ("Sr*", "100Fe/K"),

    5: ("Sr*", "K/Ca"),

    6: ("Rb*", "100Mn/Fe"),

}





@dataclass

class MeasurementInput:

    measurements: pd.DataFrame

    measurement_counts: pd.DataFrame

    qc_coefficients: pd.DataFrame | None





def _read_table(sheet):

    rows = sheet.iter_rows(values_only=True)

    raw_headers = next(rows, ())

    headers = ["" if value is None else str(value).strip() for value in raw_headers]

    while headers and not headers[-1]:

        headers.pop()

    if not headers:

        raise ValueError(f"「{sheet.title}」シートの1行目に列名がありません。")

    if any(not name for name in headers) or len(headers) != len(set(headers)):

        raise ValueError(f"「{sheet.title}」シートの1行目に空白または重複した列名があります。")

    records, row_numbers = [], []

    for row_number, values in enumerate(rows, start=2):

        values = values[:len(headers)]

        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):

            continue

        records.append(values)

        row_numbers.append(row_number)

    return pd.DataFrame(records, columns=headers), row_numbers





def read_measurements(content: bytes) -> MeasurementInput:

    try:

        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)

    except (BadZipFile, OSError, ValueError, KeyError) as exc:

        raise ValueError("Excelファイルを読み込めません。有効な.xlsxファイルを選択してください。") from exc

    try:

        if "定量値" not in workbook.sheetnames:

            raise ValueError("「定量値」シートがありません。")

        data, row_numbers = _read_table(workbook["定量値"])

        missing = [name for name in REQUIRED_COLUMNS if name not in data.columns]

        if missing:

            raise ValueError("「定量値」シートに必要な列がありません：" + "、".join(missing))

        if data.empty:

            raise ValueError("「定量値」シートに測定データがありません。")

        issues = []

        for i, row in data.iterrows():

            row_number = row_numbers[i]

            sample = row["sample"]

            if pd.isna(sample) or not str(sample).strip():

                issues.append(f"{row_number}行目：sampleが空白です")

            if pd.isna(row["測定モード"]) or not str(row["測定モード"]).strip():

                issues.append(f"{row_number}行目：測定モードが空白です")

            measured_at = row["試料測定日時"]

            if pd.isna(measured_at) or not isinstance(measured_at, (date, datetime)):

                issues.append(f"{row_number}行目：試料測定日時をExcelの日付・時刻で入力してください")

            for element in ELEMENTS:

                value = row[f"{element}_ppm"]

                if isinstance(value, bool) or not isinstance(value, Real) or not isfinite(value):

                    issues.append(f"{row_number}行目：{element}_ppmに数値がありません")

        if issues:

            remaining = len(issues) - 10

            message = "入力データを確認してください。\n" + "\n".join(issues[:10])

            if remaining > 0:

                message += f"\nほか{remaining}件"

            raise ValueError(message)

        # Normalize IDs while retaining every measurement and the original text.

        data["sample"] = data["sample"].astype(str)

        if "sample_元表記" not in data.columns:

            data.insert(1, "sample_元表記", data["sample"].copy())

        data["sample"] = data["sample"].str.replace("'", "", regex=False).str.strip()

        if data["sample"].eq("").any():

            raise ValueError("アポストロフィを除くと資料IDが空白になる測定があります。")

        counts = data.groupby("sample", sort=False).size().reset_index(name="測定回数")

        counts = counts.rename(columns={"sample": "資料ID"})

        qc = None

        if "QC補正係数" in workbook.sheetnames:

            qc, _ = _read_table(workbook["QC補正係数"])

        return MeasurementInput(data, counts, qc)

    finally:

        workbook.close()





def read_ellipse_parameters(content: bytes) -> pd.DataFrame:

    """Read only L4:T10 of AS-1 and the following group sheets."""

    try:

        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)

    except (BadZipFile, OSError, ValueError, KeyError) as exc:

        raise ValueError("判別楕円パラメーターのExcelファイルを読み込めません。") from exc

    try:

        output = []

        if "AS-1" not in workbook.sheetnames:

            raise ValueError("AS-1シートがありません。")

        start = workbook.sheetnames.index("AS-1")

        for sheet_name in workbook.sheetnames[start:]:

            sheet = workbook[sheet_name]

            block = list(sheet.iter_rows(min_row=4, max_row=10, min_col=12, max_col=20, values_only=True))

            headers = tuple("" if v is None else str(v).strip() for v in block[0])

            if headers != PARAMETER_COLUMNS:

                raise ValueError(f"{sheet.title}!L4:T4の列名を確認してください。必要な列名：" + "、".join(PARAMETER_COLUMNS))

            seen = set()

            for row_number, values in enumerate(block[1:], start=5):

                where = f"{sheet.title}シート {row_number}行目"

                record = dict(zip(PARAMETER_COLUMNS, values))

                try:

                    diagram_value = float(record["Diagram"])

                    diagram = int(diagram_value)

                except (ValueError, TypeError, OverflowError) as exc:

                    raise ValueError(f"{where}：Diagramに1〜6を入力してください。") from exc

                if diagram_value != diagram or diagram not in DIAGRAMS or isinstance(record["Diagram"], bool):

                    raise ValueError(f"{where}：Diagramに1〜6の整数を入力してください。")

                if diagram in seen:

                    raise ValueError(f"{where}：D{diagram}が重複しています。")

                seen.add(diagram)

                record["Diagram"] = diagram

                group = "" if record["Group"] is None else str(record["Group"]).strip()

                if group != sheet.title:

                    raise ValueError(f"{where}：Groupの値をシート名と一致させてください。")

                record["Group"] = group

                for name in ("X", "Y"):

                    record[name] = "" if record[name] is None else str(record[name]).strip()

                x_name, y_name = DIAGRAMS[diagram]

                if (record["X"], record["Y"]) != (x_name, y_name):

                    raise ValueError(f"{where}：D{diagram}のX、Yは{x_name}、{y_name}です。")

                for name in ("x0", "y0", "a", "b", "theta"):

                    raw = record[name]

                    try:

                        number = float(raw)

                    except (TypeError, ValueError, OverflowError) as exc:

                        raise ValueError(f"{where}：{name}に数値を入力してください。数式の場合は再計算して保存してください。") from exc

                    if isinstance(raw, bool) or not isfinite(number):

                        raise ValueError(f"{where}：{name}に有限の数値を入力してください。")

                    record[name] = number

                if record["a"] <= 0 or record["b"] <= 0:

                    raise ValueError(f"{where}：半軸長a、bは0より大きい値にしてください。")

                output.append(record)

            if seen != set(DIAGRAMS):

                raise ValueError(f"{sheet.title}シートのD1〜D6が揃っていません。")

        if not output:

            raise ValueError("判別楕円パラメーターが見つかりません。各シートのL4:T10を確認してください。")

        return pd.DataFrame(output, columns=PARAMETER_COLUMNS)

    finally:

        workbook.close()





"""Ellipse scoring, per-sample voting and selection of measurements for medians."""

from collections import Counter

from dataclasses import dataclass

from datetime import datetime

from io import BytesIO



import numpy as np

import pandas as pd

from openpyxl import Workbook

from openpyxl.styles import Alignment, Font, PatternFill

from openpyxl.utils import get_column_letter







@dataclass

class EllipseResults:

    data: pd.DataFrame

    details: pd.DataFrame

    summaries: pd.DataFrame

    medians: pd.DataFrame

    parameters: pd.DataFrame





def public_table(frame):

    return frame.loc[:, [c for c in frame.columns if not c.startswith('_')]]





def calculate_features(source):

    data = pd.DataFrame({

        '測定No': np.arange(1, len(source) + 1),

        'Sample': source['sample'].astype(str).to_numpy(),

        'Method': source['測定モード'].to_numpy(),

        'Date': source['試料測定日時'].to_numpy(),

    })

    for element in ELEMENTS:

        data[element] = source[f'{element}_ppm'].to_numpy(dtype=float)

    data['資料内測定No'] = data.groupby('Sample', sort=False).cumcount() + 1

    denominator = data[['Rb', 'Sr', 'Y', 'Zr']].sum(axis=1)

    with np.errstate(divide='ignore', invalid='ignore', over='ignore'):

        for element in ['Rb', 'Sr', 'Zr']:

            data[f'{element}*'] = 100 * data[element] / denominator

        data['K/Ca'] = data['K'] / data['Ca']

        data['100Fe/K'] = 100 * data['Fe'] / data['K']

        data['100Mn/Fe'] = 100 * data['Mn'] / data['Fe']

    features = ['Rb*', 'Sr*', 'Zr*', 'K/Ca', '100Fe/K', '100Mn/Fe']

    reasons = []

    for i, row in data.iterrows():

        problems = []

        if not np.isfinite(denominator.iloc[i]) or denominator.iloc[i] <= 0:

            problems.append('Rb+Sr+Y+Zrが0以下または計算不能')

        for element in ['Ca', 'K', 'Fe']:

            if row[element] == 0:

                problems.append(f'{element}が0のため比を計算不能')

        if not np.isfinite(row[features].to_numpy(dtype=float)).all() and not problems:

            problems.append('判別指標が計算不能')

        reasons.append('；'.join(problems))

    data['計算不能理由'] = reasons

    data['計算状態'] = np.where(data['計算不能理由'].eq(''), '計算可', '計算不能')

    # Unavailable values must not be exported as infinity or used for scoring.

    data[features] = data[features].replace([np.inf, -np.inf], np.nan)

    return data





def normalized_distance(x, y, parameter):

    theta = np.deg2rad(parameter['theta'])

    dx, dy = x - parameter['x0'], y - parameter['y0']

    with np.errstate(over='ignore', invalid='ignore', divide='ignore'):

        return np.hypot((dx * np.cos(theta) + dy * np.sin(theta)) / parameter['a'],

                        (-dx * np.sin(theta) + dy * np.cos(theta)) / parameter['b'])





def score_distances(distances, computable):

    """Rows are measurements; columns are D1...D6. Boundaries are inclusive."""

    finite = np.isfinite(distances).all(axis=1)

    inside = distances <= 1

    mandatory = inside[:, :3].all(axis=1)

    auxiliary = (distances[:, 3:] <= np.array([3., 3., 5.])).all(axis=1)

    valid = np.asarray(computable) & finite & mandatory & auxiliary

    scores = np.where(valid, 7 + inside[:, 3:].sum(axis=1), np.nan)

    return inside, valid, scores





def classify_measurements(data, parameters):

    frames = []

    for group, group_parameters in parameters.groupby('Group', sort=False):

        p = group_parameters.set_index('Diagram')

        if len(p) != 6 or set(p.index) != set(DIAGRAMS):

            raise ValueError(f'{group}のD1〜D6が揃っていません。')

        distances = np.column_stack([

            normalized_distance(data[x].to_numpy(), data[y].to_numpy(), p.loc[d])

            for d, (x, y) in DIAGRAMS.items()

        ])

        computable = data['計算状態'].eq('計算可').to_numpy()

        inside, valid, scores = score_distances(distances, computable)

        frame = data[['測定No', 'Sample', '資料内測定No']].copy()

        frame['化学組成グループ'] = group

        for d in range(1, 7):

            # A nonfinite distance cannot support an inside/outside decision.

            frame[f'D{d}規格化距離'] = np.where(np.isfinite(distances[:, d-1]), distances[:, d-1], np.nan)

            frame[f'D{d}判定'] = np.where(~computable | ~np.isfinite(distances[:, d-1]),

                                        '計算不能', np.where(inside[:, d-1], '楕円内', '楕円外'))

        frame['有効性'] = np.where(valid, '有効', '無効')

        frame['得点'] = scores

        reasons = []

        for i in range(len(data)):

            if not computable[i]:

                reasons.append(data.iloc[i]['計算不能理由'])

            elif not np.isfinite(distances[i]).all():

                reasons.append('規格化距離が計算不能')

            else:

                outside = [f'D{d}' for d in range(1, 4) if not inside[i, d-1]]

                exceeded = [f'D{d}' for d, limit in [(4, 3), (5, 3), (6, 5)] if distances[i, d-1] > limit]

                parts = []

                if outside:

                    parts.append('・'.join(outside) + 'が楕円外')

                if exceeded:

                    parts.append('・'.join(exceeded) + 'の規格化距離上限超過')

                reasons.append('；'.join(parts))

        frame['無効理由'] = reasons

        frames.append(frame)

    if not frames:

        raise ValueError('判別楕円パラメーターがありません。')

    details = pd.concat(frames, ignore_index=True).sort_values('測定No', kind='stable').reset_index(drop=True)

    valid_details = details[details['有効性'].eq('有効')]

    best = {}

    for measurement_no, part in valid_details.groupby('測定No', sort=False):

        high_score = int(part['得点'].max())

        groups = tuple(part.loc[part['得点'].eq(high_score), '化学組成グループ'])

        best[measurement_no] = (high_score, groups)

    decisions = data.copy()

    decisions['_Groups'] = [best.get(no, (None, ()))[1] for no in decisions['測定No']]

    decisions['判別結果'] = decisions['_Groups'].map(lambda groups: '・'.join(groups) if groups else '判定不能')

    decisions['得点'] = [best.get(no, (np.nan, ()))[0] for no in decisions['測定No']]

    decisions['有効性'] = np.where(decisions['_Groups'].map(bool), '有効', '判定不能')

    return decisions, details





def aggregate_measurements(decisions, group_order):

    rows = []

    for sample, part in decisions.groupby('Sample', sort=False):

        valid = part[part['有効性'].eq('有効')]

        final_groups, high_score = (), np.nan

        if not valid.empty:

            high_score = int(valid['得点'].max())

            highest = valid[valid['得点'].eq(high_score)]

            max_counts = Counter(g for groups in highest['_Groups'] for g in groups)

            candidates = {g for g, count in max_counts.items() if count == max(max_counts.values())}

            if len(candidates) > 1:

                all_counts = Counter(g for groups in valid['_Groups'] for g in groups)

                max_all = max(all_counts[g] for g in candidates)

                candidates = {g for g in candidates if all_counts[g] == max_all}

            final_groups = tuple(g for g in group_order if g in candidates)

        rows.append({'Sample': sample,

                     '最終判別結果': '・'.join(final_groups) if final_groups else '判定不能',

                     '得点': high_score, '測定データ数': len(part), '有効測定数': len(valid),

                     '_Groups': final_groups})

    return pd.DataFrame(rows)





def select_medians(decisions, summaries):

    selected = pd.Series(False, index=decisions.index)

    medians, counts = [], []

    for _, summary in summaries.iterrows():

        mask = decisions['Sample'].eq(summary['Sample'])

        if summary['最終判別結果'] != '判定不能':

            final_groups = set(summary['_Groups'])

            mask &= decisions['有効性'].eq('有効') & decisions['得点'].eq(summary['得点'])

            mask &= decisions['_Groups'].map(lambda groups: bool(final_groups.intersection(groups)))

        rows = decisions.loc[mask]

        if rows.empty:

            raise ValueError(f"{summary['Sample']}の中央値に使用できる測定データがありません。")

        selected |= mask

        counts.append(len(rows))

        median = {'Sample': summary['Sample']}

        median.update(rows[list(ELEMENTS)].median().to_dict())

        median.update({'主要判定結果': summary['最終判別結果'], '得点': summary['得点'], '中央値使用数': len(rows)})

        medians.append(median)

    decisions = decisions.copy()

    decisions['中央値使用'] = np.where(selected, '使用', '除外')

    summaries = summaries.copy()

    summaries['中央値使用数'] = counts

    return decisions, summaries, pd.DataFrame(medians)





def run_ellipse_analysis(inputs: MeasurementInput, parameters: pd.DataFrame) -> EllipseResults:

    features = calculate_features(inputs.measurements)

    decisions, details = classify_measurements(features, parameters)

    summaries = aggregate_measurements(decisions, list(parameters['Group'].unique()))

    decisions, summaries, medians = select_medians(decisions, summaries)

    return EllipseResults(decisions, details, summaries, medians, parameters.copy())





def _write_table(sheet, frame, start_row=1, start_col=1):

    """Store computed values, preserving literal sample IDs as Excel text."""

    frame = public_table(frame)

    for offset, values in enumerate([list(frame.columns)] + list(frame.itertuples(index=False, name=None))):

        for j, value in enumerate(values):

            if pd.isna(value):

                value = None

            elif isinstance(value, np.generic):

                value = value.item()

            elif isinstance(value, pd.Timestamp):

                value = value.to_pydatetime()

            cell = sheet.cell(start_row + offset, start_col + j, value)

            if isinstance(value, str):

                cell.data_type = 's'

    end_col = start_col + len(frame.columns) - 1

    for cell in sheet.iter_rows(min_row=start_row, max_row=start_row, min_col=start_col, max_col=end_col):

        for c in cell:

            c.fill = PatternFill('solid', fgColor='365F91')

            c.font = Font(color='FFFFFF', bold=True)

            c.alignment = Alignment(horizontal='center', vertical='center')

    for j, column in enumerate(frame.columns, start=start_col):

        width = max(13, min(28, len(str(column)) * 1.8 + 3))

        sheet.column_dimensions[get_column_letter(j)].width = width

        for cells in sheet.iter_rows(min_row=start_row + 1, min_col=j, max_col=j, max_row=start_row + len(frame)):

            cell = cells[0]

            cell.alignment = Alignment(horizontal='center', vertical='center')

            if isinstance(cell.value, float):

                cell.number_format = '0' if column == '得点' else '0.0000'

            if isinstance(cell.value, datetime):

                cell.number_format = 'yyyy-mm-dd hh:mm'

    sheet.freeze_panes = sheet.cell(start_row + 1, start_col + 1).coordinate

    sheet.auto_filter.ref = f'{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{start_row + len(frame)}'

    sheet.sheet_view.showGridLines = False





def export_results(results: EllipseResults, inputs: MeasurementInput) -> bytes:

    workbook = Workbook()

    workbook.remove(workbook.active)

    _write_table(workbook.create_sheet('RESULTS'), public_table(results.summaries), start_col=2)

    median_ws = workbook.create_sheet('Median')

    _write_table(median_ws, results.medians, start_row=3, start_col=2)

    for c in range(3, 13):

        median_ws.cell(1, c, 'Median')

        median_ws.cell(2, c, 'ppm')

    # Keep the element and ratio columns in the same positions as the Excel template.

    columns = ['測定No', 'Sample', 'Method', 'Date'] + list(ELEMENTS)

    columns += ['Rb*', 'Sr*', 'Zr*', 'K/Ca', '100Fe/K', '100Mn/Fe']

    data = results.data[columns].copy()

    data['資料ID'] = results.data['Sample']

    data['資料内測定No'] = results.data['資料内測定No']

    for name in ['判別結果', '得点', '有効性', '中央値使用', '計算状態', '計算不能理由']:

        data[name] = results.data[name]

    _write_table(workbook.create_sheet('DATA'), data, start_row=2)

    _write_table(workbook.create_sheet('判別詳細'), results.details)

    _write_table(workbook.create_sheet('定量値'), inputs.measurements)

    if inputs.qc_coefficients is not None:

        _write_table(workbook.create_sheet('QC補正係数'), inputs.qc_coefficients)

    for group, parameters in results.parameters.groupby('Group', sort=False):

        if group in workbook.sheetnames:

            raise ValueError(f'化学組成グループ名「{group}」が出力シート名と重複しています。')

        sheet = workbook.create_sheet(group)

        sheet['L3'] = group

        _write_table(sheet, parameters[list(PARAMETER_COLUMNS)], start_row=4, start_col=12)

        for cells in sheet.iter_rows(min_row=5, max_row=10, min_col=16, max_col=20):

            for cell in cells:

                cell.number_format = '0.00000000'

    buffer = BytesIO()

    workbook.save(buffer)

    return buffer.getvalue()
