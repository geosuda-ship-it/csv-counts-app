#!/usr/bin/env python3
"""クラスター楕円と化学組成グループの判別楕円を比較する。

Google Colabでは、実行すると2つのExcelファイルのアップロード画面が順に開く。
通常のPCではファイル選択画面を使う。コマンドラインから2ファイルを指定してもよい。

出力: クラスター区分ファイル名に基づく「*_判定結果.xlsx」
  - 候補一覧: Cluster別のD1-D3共通候補とD1-D6共通候補
  - 試料別判定: Sample別のクラスターと候補
  - 重なり結果: Diagramごとの重複率と判定
  - Median: クラスター区分ファイルの同名シート（計算済み値を優先）
"""

from __future__ import annotations

import argparse
import math
import re
import sys
from copy import copy, deepcopy
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


DIAGRAMS = {
    1: ("Rb*", "Sr*"),
    2: ("Rb*", "Zr*"),
    3: ("Sr*", "Zr*"),
    4: ("Sr*", "100Fe/K"),
    5: ("Sr*", "K/Ca"),
    6: ("Rb*", "100Mn/Fe"),
}
SCRIPT_VERSION = "2026-10-24-v5"
MEDIAN_SHEET_NAME = "Median"


def natural_sample_key(value) -> tuple:
    """Sample-2をSample-10より前に並べるための自然順キー。"""
    text = str(value or "").strip()
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", text)
        if part
    )


@dataclass(frozen=True)
class Ellipse:
    name: str
    diagram: int
    x_name: str
    y_name: str
    x0: float
    y0: float
    a: float
    b: float
    theta_deg: float
    priority: float = 1.0

    @property
    def area(self) -> float:
        return math.pi * self.a * self.b

    @property
    def quadratic(self) -> tuple[float, float, float]:
        """A*dx^2 + B*dx*dy + C*dy^2 <= 1 のA, B, Cを返す。"""
        theta = math.radians(self.theta_deg)
        cos_t, sin_t = math.cos(theta), math.sin(theta)
        inv_a2, inv_b2 = 1.0 / self.a**2, 1.0 / self.b**2
        return (
            cos_t**2 * inv_a2 + sin_t**2 * inv_b2,
            2.0 * cos_t * sin_t * (inv_a2 - inv_b2),
            sin_t**2 * inv_a2 + cos_t**2 * inv_b2,
        )

    @property
    def x_bounds(self) -> tuple[float, float]:
        theta = math.radians(self.theta_deg)
        half_width = math.sqrt(
            (self.a * math.cos(theta)) ** 2 + (self.b * math.sin(theta)) ** 2
        )
        return self.x0 - half_width, self.x0 + half_width


def number(value, location: str) -> float:
    if value is None or isinstance(value, str):
        raise ValueError(f"数値がありません: {location} ({value!r})")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"有限の数値ではありません: {location} ({value!r})")
    return result


def read_cluster_ellipses(path: Path) -> list[Ellipse]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "散布図" not in workbook.sheetnames:
        raise ValueError(f"{path.name} に「散布図」シートがありません")
    sheet = workbook["散布図"]
    result: list[Ellipse] = []
    for row in range(5, 59):
        name = sheet.cell(row, 16).value       # P
        diagram = sheet.cell(row, 13).value    # M
        values = [sheet.cell(row, col).value for col in range(17, 22)]  # Q:U
        if not name or diagram not in DIAGRAMS or any(v is None for v in values):
            continue
        diagram = int(diagram)
        x_name, y_name = sheet.cell(row, 14).value, sheet.cell(row, 15).value
        if (x_name, y_name) != DIAGRAMS[diagram]:
            raise ValueError(f"散布図!N{row}:O{row} の軸とD{diagram}の定義が一致しません")
        x0, y0, a, b, theta = (
            number(v, f"散布図!{col}{row}") for v, col in zip(values, "QRSTU")
        )
        if a <= 0 or b <= 0:
            raise ValueError(f"散布図!S{row}:T{row} の軸半径は正数である必要があります")
        result.append(Ellipse(str(name), diagram, x_name, y_name, x0, y0, a, b, theta))
    workbook.close()
    if not result:
        raise ValueError("「散布図」シートから有効なクラスター楕円を取得できませんでした")
    return result


def read_group_ellipses(path: Path, first_sheet: str = "AS-1") -> list[Ellipse]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if first_sheet not in workbook.sheetnames:
        raise ValueError(f"{path.name} にシート {first_sheet!r} がありません")
    start = workbook.sheetnames.index(first_sheet)
    result: list[Ellipse] = []
    for sheet_name in workbook.sheetnames[start:]:
        sheet = workbook[sheet_name]
        for row in range(5, 11):
            diagram = sheet.cell(row, 12).value  # L
            name = sheet.cell(row, 13).value      # M
            if diagram not in DIAGRAMS or not name:
                continue
            diagram = int(diagram)
            x_name, y_name = sheet.cell(row, 14).value, sheet.cell(row, 15).value
            if (x_name, y_name) != DIAGRAMS[diagram]:
                raise ValueError(f"{sheet_name}!N{row}:O{row} の軸とD{diagram}の定義が一致しません")
            values = [sheet.cell(row, col).value for col in range(16, 22)]  # P:U
            x0, y0, a, b, theta, priority = (
                number(v, f"{sheet_name}!{col}{row}")
                for v, col in zip(values, "PQRSTU")
            )
            if a <= 0 or b <= 0:
                raise ValueError(f"{sheet_name}!R{row}:S{row} の軸半径は正数である必要があります")
            result.append(
                Ellipse(str(name), diagram, x_name, y_name, x0, y0, a, b, theta, priority)
            )
    workbook.close()
    if not result:
        raise ValueError("判別楕円シートから有効な楕円を取得できませんでした")
    return result


def vertical_interval(ellipse: Ellipse, x: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    qa, qb, qc = ellipse.quadratic
    dx = x - ellipse.x0
    linear = qb * dx
    discriminant = linear**2 - 4.0 * qc * (qa * dx**2 - 1.0)
    valid = discriminant >= -1e-12
    root = np.sqrt(np.maximum(discriminant, 0.0))
    lower = (-linear - root) / (2.0 * qc)
    upper = (-linear + root) / (2.0 * qc)
    return ellipse.y0 + lower, ellipse.y0 + upper, valid


def intersection_area(
    first: Ellipse, second: Ellipse, nodes: np.ndarray, weights: np.ndarray
) -> float:
    lower_x = max(first.x_bounds[0], second.x_bounds[0])
    upper_x = min(first.x_bounds[1], second.x_bounds[1])
    if upper_x <= lower_x:
        return 0.0
    x = (upper_x - lower_x) * 0.5 * nodes + (upper_x + lower_x) * 0.5
    low_1, high_1, valid_1 = vertical_interval(first, x)
    low_2, high_2, valid_2 = vertical_interval(second, x)
    height = np.maximum(0.0, np.minimum(high_1, high_2) - np.maximum(low_1, low_2))
    height[~(valid_1 & valid_2)] = 0.0
    area = (upper_x - lower_x) * 0.5 * float(np.dot(weights, height))
    return max(0.0, min(area, first.area, second.area))


def classify(cluster_pct: float, group_pct: float, area: float, tolerance_pct: float) -> str:
    if cluster_pct >= tolerance_pct and group_pct >= tolerance_pct:
        return "ほぼ一致"
    if cluster_pct >= tolerance_pct:
        return "クラスター楕円が判別楕円内"
    if group_pct >= tolerance_pct:
        return "判別楕円がクラスター楕円内"
    if area > 1e-10:
        return "部分重複"
    return "非重複"


def calculate_overlaps(
    clusters: list[Ellipse],
    groups: list[Ellipse],
    quadrature_points: int,
    tolerance_pct: float,
) -> list[dict]:
    nodes, weights = np.polynomial.legendre.leggauss(quadrature_points)
    group_lookup = {(ellipse.name, ellipse.diagram): ellipse for ellipse in groups}
    cluster_lookup = {(ellipse.name, ellipse.diagram): ellipse for ellipse in clusters}
    group_names = list(dict.fromkeys(ellipse.name for ellipse in groups))
    cluster_names = list(dict.fromkeys(ellipse.name for ellipse in clusters))
    detail: list[dict] = []

    for cluster_name in cluster_names:
        for group_name in group_names:
            for diagram in DIAGRAMS:
                cluster = cluster_lookup.get((cluster_name, diagram))
                group = group_lookup.get((group_name, diagram))
                if cluster is None or group is None:
                    continue
                area = intersection_area(cluster, group, nodes, weights)
                cluster_pct = 100.0 * area / cluster.area
                group_pct = 100.0 * area / group.area
                union = cluster.area + group.area - area
                iou_pct = 100.0 * area / union if union > 0 else 0.0

                # 非重複と、Excel上で全指標が0.0%に丸められる極小重複は除外する。
                if area <= 1e-10 or max(cluster_pct, group_pct, iou_pct) < 0.05:
                    continue
                detail.append(
                    {
                        "Cluster": cluster_name,
                        "Group": group_name,
                        "Diagram": f"D{diagram}",
                        "Cluster包含率": cluster_pct / 100.0,
                        "Group占有率": group_pct / 100.0,
                        "IoU": iou_pct / 100.0,
                        "判定": classify(cluster_pct, group_pct, area, tolerance_pct),
                    }
                )
    return detail


def summarize_candidates(detail: list[dict], cluster_names: list[str]) -> list[dict]:
    coverage: dict[tuple[str, str], set[str]] = {}
    iou_lookup: dict[tuple[str, str, str], float] = {}
    for row in detail:
        key = (row["Cluster"], row["Group"])
        coverage.setdefault(key, set()).add(row["Diagram"])
        iou_lookup[(row["Cluster"], row["Group"], row["Diagram"])] = row["IoU"]

    required_3 = {"D1", "D2", "D3"}
    required_6 = {f"D{diagram}" for diagram in DIAGRAMS}

    def ranked_values(cluster: str, required: set[str]) -> list[tuple[float, float, str]]:
        candidates: list[tuple[float, float, str]] = []
        for candidate_cluster, group in coverage:
            if candidate_cluster != cluster or not required <= coverage[(cluster, group)]:
                continue
            values = [iou_lookup[(cluster, group, diagram)] for diagram in sorted(required)]
            average_iou = sum(values) / len(values)
            minimum_iou = min(values)
            candidates.append((average_iou, minimum_iou, group))

        # 平均IoUの高い順。平均が同じ場合は最小IoUの高い方を優先する。
        candidates.sort(key=lambda item: (-item[0], -item[1], item[2]))
        return candidates

    def full_labels(candidates: list[tuple[float, float, str]]) -> str:
        if not candidates:
            return "該当なし"
        return "・".join(
            f"{group}（平均IoU {average_iou:.1%}、最小IoU {minimum_iou:.1%}）"
            for average_iou, minimum_iou, group in candidates
        )

    def report_labels(candidates: list[tuple[float, float, str]]) -> str:
        if not candidates:
            return "該当なし"
        return "・".join(
            f"{group}（{average_iou:.1%}）"
            for average_iou, _minimum_iou, group in candidates
        )

    summary: list[dict] = []
    for cluster in cluster_names:
        common_3 = ranked_values(cluster, required_3)
        common_6 = ranked_values(cluster, required_6)
        summary.append(
            {
                "Cluster": cluster,
                "3図共通候補（D1-D3）": full_labels(common_3),
                "6図共通候補（D1-D6）": full_labels(common_6),
                "3図報告表記": report_labels(common_3),
                "6図報告表記": report_labels(common_6),
            }
        )
    return summary


def read_sample_assignments(path: Path) -> list[dict]:
    """DATAシートからSample名と最終的なクラスター区分を取得する。"""
    workbook = load_workbook(path, data_only=True, read_only=True)
    if "DATA" not in workbook.sheetnames:
        raise ValueError(f"{path.name} に「DATA」シートがありません")
    sheet = workbook["DATA"]
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]

    try:
        sample_column = headers.index("Sample") + 1
    except ValueError as exc:
        raise ValueError("DATAシートに「Sample」列がありません") from exc

    if "区分結果" in headers:
        cluster_column = headers.index("区分結果") + 1
    elif "一次区分" in headers:
        cluster_column = headers.index("一次区分") + 1
    else:
        raise ValueError("DATAシートに「区分結果」または「一次区分」列がありません")

    result: list[dict] = []
    for row in range(2, sheet.max_row + 1):
        sample = sheet.cell(row, sample_column).value
        cluster = sheet.cell(row, cluster_column).value
        if sample in (None, ""):
            continue
        result.append(
            {
                "Sample": str(sample),
                "Cluster": "" if cluster is None else str(cluster),
            }
        )
    workbook.close()
    result.sort(key=lambda row: natural_sample_key(row["Sample"]))
    return result


def require_median_sheet(path: Path) -> None:
    """入力ブックに必須のMedianシートがあることを早い段階で確認する。"""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if MEDIAN_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"{path.name} に必須の「{MEDIAN_SHEET_NAME}」シートがありません。"
                "シート名（大文字・小文字を含む）を確認してください"
            )
    finally:
        workbook.close()


def _sort_and_filter_median_sheet(
    sheet,
    sample_assignments: list[dict],
) -> None:
    """Medianを資料別判定と同じSample順に並べ、フィルタを設定する。"""
    header_row = 3
    headers = {
        str(sheet.cell(header_row, column).value).strip(): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(header_row, column).value not in (None, "")
    }
    sample_column = headers.get("Sample")
    if sample_column is None:
        raise ValueError("Medianシートの3行目に「Sample」列がありません")

    order = {
        str(record["Sample"]).strip(): index
        for index, record in enumerate(sample_assignments)
    }
    data_rows = []
    for row_number in range(header_row + 1, sheet.max_row + 1):
        values = [
            sheet.cell(row_number, column).value
            for column in range(1, sheet.max_column + 1)
        ]
        sample = values[sample_column - 1]
        if sample in (None, ""):
            continue
        sample_text = str(sample).strip()
        data_rows.append((sample_text, values))

    data_rows.sort(
        key=lambda item: (
            0,
            order[item[0]],
        )
        if item[0] in order
        else (
            1,
            natural_sample_key(item[0]),
        )
    )

    for target_row, (_sample, values) in enumerate(
        data_rows,
        start=header_row + 1,
    ):
        for column, value in enumerate(values, start=1):
            sheet.cell(target_row, column).value = value

    # 元の末尾に空白行が混在していた場合は、並べ替え後の余分な値を消す。
    first_unused_row = header_row + 1 + len(data_rows)
    for row_number in range(first_unused_row, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            sheet.cell(row_number, column).value = None

    first_header_column = min(headers.values())
    last_header_column = max(headers.values())
    last_data_row = max(header_row, header_row + len(data_rows))
    sheet.auto_filter.ref = (
        f"{get_column_letter(first_header_column)}{header_row}:"
        f"{get_column_letter(last_header_column)}{last_data_row}"
    )


def copy_median_sheet(
    source_path: Path,
    destination_workbook: Workbook,
    sample_assignments: list[dict],
) -> None:
    """Medianを値中心で複製し、openpyxlで扱える書式・寸法を引き継ぐ。

    Excelが保存した数式キャッシュ（計算済み値）が存在するときはその値を採用する。
    キャッシュが空の数式セルだけは、情報を失わないよう元の数式を残す。
    """
    source_formula = load_workbook(source_path, data_only=False, read_only=False)
    source_values = load_workbook(source_path, data_only=True, read_only=False)
    try:
        if MEDIAN_SHEET_NAME not in source_formula.sheetnames:
            raise ValueError(
                f"{source_path.name} に必須の「{MEDIAN_SHEET_NAME}」シートがありません。"
                "シート名（大文字・小文字を含む）を確認してください"
            )

        formula_sheet = source_formula[MEDIAN_SHEET_NAME]
        value_sheet = source_values[MEDIAN_SHEET_NAME]
        target = destination_workbook.create_sheet(MEDIAN_SHEET_NAME)

        # シート全体の表示・印刷設定。
        target.sheet_format = copy(formula_sheet.sheet_format)
        target.sheet_properties = copy(formula_sheet.sheet_properties)
        target.views = copy(formula_sheet.views)
        target.freeze_panes = formula_sheet.freeze_panes
        target.sheet_state = formula_sheet.sheet_state
        target.page_margins = copy(formula_sheet.page_margins)
        target.page_setup = copy(formula_sheet.page_setup)
        target.print_options = copy(formula_sheet.print_options)
        target.sheet_properties.pageSetUpPr = copy(formula_sheet.sheet_properties.pageSetUpPr)
        target.auto_filter = copy(formula_sheet.auto_filter)
        target.protection = copy(formula_sheet.protection)
        target.print_area = formula_sheet.print_area
        target.print_title_cols = formula_sheet.print_title_cols
        target.print_title_rows = formula_sheet.print_title_rows

        for key, dimension in formula_sheet.column_dimensions.items():
            target.column_dimensions[key] = copy(dimension)
        for key, dimension in formula_sheet.row_dimensions.items():
            target.row_dimensions[key] = copy(dimension)

        for row in formula_sheet.iter_rows():
            for source_cell in row:
                target_cell = target[source_cell.coordinate]
                calculated = value_sheet[source_cell.coordinate].value
                # 通常セルは元値を採用。数式セルはキャッシュがあれば値を優先する。
                if source_cell.data_type == "f" and calculated is not None:
                    target_cell.value = calculated
                else:
                    target_cell.value = source_cell.value
                if source_cell.has_style:
                    target_cell._style = copy(source_cell._style)
                if source_cell.number_format:
                    target_cell.number_format = source_cell.number_format
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.protection = copy(source_cell.protection)
                if source_cell.hyperlink:
                    target_cell._hyperlink = copy(source_cell.hyperlink)
                if source_cell.comment:
                    target_cell.comment = copy(source_cell.comment)

        for merged_range in formula_sheet.merged_cells.ranges:
            target.merge_cells(str(merged_range))

        # 条件付き書式・入力規則は参照式も含め、可能な範囲でそのまま保持する。
        target.conditional_formatting = deepcopy(formula_sheet.conditional_formatting)
        target.data_validations = deepcopy(formula_sheet.data_validations)
        _sort_and_filter_median_sheet(target, sample_assignments)
    finally:
        source_formula.close()
        source_values.close()


def style_header(cells, fill_color: str = "1F4E78") -> None:
    fill = PatternFill("solid", fgColor=fill_color)
    for cell in cells:
        cell.fill = fill
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_results(
    path: Path,
    cluster_path: Path,
    detail: list[dict],
    summary: list[dict],
    sample_assignments: list[dict],
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "候補一覧"
    sample_sheet = workbook.create_sheet("試料別判定")
    detail_sheet = workbook.create_sheet("重なり結果")

    summary_sheet.sheet_view.showGridLines = False
    summary_sheet["A2"] = "Cluster別 重複Group候補"
    summary_sheet["A2"].font = Font(name="Arial", size=14, bold=True)
    summary_headers = ["Cluster", "3図共通候補（D1-D3）", "6図共通候補（D1-D6）"]
    summary_sheet.append([])
    summary_sheet.append(summary_headers)
    for row in summary:
        summary_sheet.append([row[header] for header in summary_headers])
    style_header(summary_sheet[4])
    summary_sheet.freeze_panes = "A5"
    summary_sheet.column_dimensions["A"].width = 13
    summary_sheet.column_dimensions["B"].width = 62
    summary_sheet.column_dimensions["C"].width = 62
    thin_gray = Side(style="thin", color="D9D9D9")
    for row in summary_sheet.iter_rows(min_row=5, max_row=4 + len(summary), max_col=3):
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=(cell.column == 1))
            cell.alignment = Alignment(
                horizontal="center" if cell.column == 1 else "left",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin_gray)
        row[0].fill = PatternFill("solid", fgColor="D9EAF7")
    for excel_row, record in enumerate(summary, start=5):
        longest = max(
            len(record["3図共通候補（D1-D3）"]),
            len(record["6図共通候補（D1-D6）"]),
        )
        estimated_lines = max(1, math.ceil(longest / 55))
        summary_sheet.row_dimensions[excel_row].height = max(24, 15 * estimated_lines)
    note_row = 6 + len(summary)
    summary_sheet.cell(note_row, 1, "判定方法")
    summary_sheet.cell(
        note_row,
        2,
        "各Cluster・Groupについて、指定したすべてのDiagramに重複レコードが存在する場合に候補としました。",
    )
    summary_sheet.cell(note_row + 1, 1, "補足")
    summary_sheet.cell(
        note_row + 1,
        2,
        "3図共通候補には、D4-D6のいずれかで重ならないGroupも残しています。",
    )
    for row in summary_sheet.iter_rows(min_row=note_row, max_row=note_row + 1, max_col=3):
        for cell in row:
            cell.font = Font(name="Arial", size=9, italic=(cell.column > 1), bold=(cell.column == 1), color="595959")
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    summary_lookup = {row["Cluster"]: row for row in summary}
    sample_headers = [
        "Sample",
        "Cluster",
        "3図共通候補（平均IoU）",
        "6図共通候補（平均IoU）",
    ]
    sample_sheet.append(sample_headers)
    for assignment in sample_assignments:
        cluster = assignment["Cluster"]
        cluster_summary = summary_lookup.get(cluster)
        if cluster_summary is not None:
            common_3 = cluster_summary["3図報告表記"]
            common_6 = cluster_summary["6図報告表記"]
        elif cluster:
            common_3 = cluster
            common_6 = cluster
        else:
            common_3 = ""
            common_6 = ""
        sample_sheet.append(
            [assignment["Sample"], cluster, common_3, common_6]
        )
    style_header(sample_sheet[1], "2F75B5")
    sample_sheet.freeze_panes = "A2"
    sample_sheet.sheet_view.showGridLines = False
    sample_sheet.column_dimensions["A"].width = 18
    sample_sheet.column_dimensions["B"].width = 12
    sample_sheet.column_dimensions["C"].width = 55
    sample_sheet.column_dimensions["D"].width = 55
    for row in sample_sheet.iter_rows(min_row=2, max_col=4):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(
                horizontal="center" if cell.column == 2 else "left",
                vertical="center",
                wrap_text=cell.column >= 3,
            )
    if sample_assignments:
        sample_table = Table(displayName="SampleJudgmentResults", ref=sample_sheet.dimensions)
        sample_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sample_sheet.add_table(sample_table)

    detail_headers = [
        "Cluster", "Group", "Diagram", "Cluster包含率", "Group占有率", "IoU", "判定"
    ]
    detail_sheet.append(detail_headers)
    for row in detail:
        detail_sheet.append([row[header] for header in detail_headers])
    style_header(detail_sheet[1], "4472C4")
    detail_sheet.freeze_panes = "A2"
    detail_sheet.sheet_view.showGridLines = False
    for row in detail_sheet.iter_rows(min_row=2, max_col=7):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in detail_sheet.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "0.0%"
    for column, width in {
        "A": 12, "B": 14, "C": 11, "D": 16, "E": 14, "F": 11, "G": 30
    }.items():
        detail_sheet.column_dimensions[column].width = width
    if detail:
        table = Table(displayName="EllipseOverlapResults", ref=detail_sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        detail_sheet.add_table(table)
    copy_median_sheet(cluster_path, workbook, sample_assignments)
    workbook.save(path)


def workbook_role(path: Path) -> str | None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_names = set(workbook.sheetnames)
    workbook.close()
    if "散布図" in sheet_names:
        return "cluster"
    if "AS-1" in sheet_names:
        return "group"
    return None


def assign_workbooks(paths: list[Path]) -> tuple[Path, Path]:
    cluster_path = None
    group_path = None
    for path in paths:
        role = workbook_role(path)
        if role == "cluster":
            if cluster_path is not None:
                raise ValueError("「散布図」シートを持つファイルが複数あります")
            cluster_path = path
        elif role == "group":
            if group_path is not None:
                raise ValueError("「AS-1」シートを持つファイルが複数あります")
            group_path = path
    if cluster_path is None:
        raise ValueError("「散布図」シートを持つクラスター区分ファイルが見つかりません")
    if group_path is None:
        raise ValueError("「AS-1」シートを持つ判別楕円パラメータファイルが見つかりません")
    return cluster_path, group_path


def running_in_colab() -> bool:
    return "google.colab" in sys.modules


def select_files_interactively() -> tuple[Path, Path, Path, bool]:
    if running_in_colab():
        from google.colab import files

        print(f"【{SCRIPT_VERSION}】1/2：クラスター区分_解析用ファイルを選択してください。")
        first = files.upload()
        if len(first) != 1:
            raise ValueError("Excelファイルを1つだけアップロードしてください")
        print(f"【{SCRIPT_VERSION}】2/2：判別楕円パラメータファイルを選択してください。")
        second = files.upload()
        if len(second) != 1:
            raise ValueError("Excelファイルを1つだけアップロードしてください")
        paths = [Path(next(iter(first))), Path(next(iter(second)))]
        cluster_path, group_path = assign_workbooks(paths)
        return cluster_path, group_path, Path.cwd(), True

    try:
        from tkinter import Tk, filedialog
    except ImportError as exc:
        raise RuntimeError(
            "ファイル選択画面を開けません。コマンドライン引数で2ファイルを指定してください。"
        ) from exc
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    selected: list[Path] = []
    try:
        for title in ("クラスター区分解析用ファイルを選択", "判別楕円パラメータファイルを選択"):
            name = filedialog.askopenfilename(
                title=title,
                filetypes=[("Excelファイル", "*.xlsx"), ("すべてのファイル", "*.*")],
            )
            if not name:
                raise RuntimeError("ファイル選択がキャンセルされました")
            selected.append(Path(name))
    finally:
        root.destroy()
    cluster_path, group_path = assign_workbooks(selected)
    return cluster_path, group_path, cluster_path.parent, False


def analyze(
    cluster_path: Path,
    group_path: Path,
    output_dir: Path,
    quadrature_points: int,
    tolerance_pct: float,
) -> Path:
    require_median_sheet(cluster_path)
    clusters = read_cluster_ellipses(cluster_path)
    groups = read_group_ellipses(group_path)
    sample_assignments = read_sample_assignments(cluster_path)
    detail = calculate_overlaps(clusters, groups, quadrature_points, tolerance_pct)
    cluster_names = list(dict.fromkeys(ellipse.name for ellipse in clusters))
    summary = summarize_candidates(detail, cluster_names)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_name = cluster_path.stem.removesuffix("_ver1") + "_判定結果.xlsx"
    output_path = output_dir / output_name
    write_results(output_path, cluster_path, detail, summary, sample_assignments)

    print(f"Clusters: {len(cluster_names)} ({', '.join(cluster_names)})")
    print(f"Chemical groups: {len(set(ellipse.name for ellipse in groups))}")
    print(f"Overlap records: {len(detail)}")
    print(f"Sample judgments: {len(sample_assignments)}")
    print("\nCluster別候補")
    for row in summary:
        print(
            f"{row['Cluster']}: "
            f"D1-D3 = {row['3図共通候補（D1-D3）']} / "
            f"D1-D6 = {row['6図共通候補（D1-D6）']}"
        )
    print(f"\nOutput: {output_path}")
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("cluster_workbook", type=Path, nargs="?")
    parser.add_argument("group_workbook", type=Path, nargs="?")
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--quadrature-points", type=int, default=512)
    parser.add_argument("--containment-tolerance", type=float, default=99.5)
    if "ipykernel" in sys.modules or running_in_colab():
        args = parser.parse_args([])
    else:
        args = parser.parse_args()

    if args.quadrature_points < 64:
        parser.error("--quadrature-points は64以上にしてください")
    if (args.cluster_workbook is None) != (args.group_workbook is None):
        parser.error("Excelファイルは2つとも指定するか、両方とも省略してください")

    should_download = False
    if args.cluster_workbook is None:
        cluster_path, group_path, default_output_dir, should_download = select_files_interactively()
    else:
        cluster_path, group_path = assign_workbooks(
            [args.cluster_workbook, args.group_workbook]
        )
        default_output_dir = Path.cwd()
    output_path = analyze(
        cluster_path,
        group_path,
        args.output_dir or default_output_dir,
        args.quadrature_points,
        args.containment_tolerance,
    )
    if should_download:
        from google.colab import files

        files.download(str(output_path))


if __name__ == "__main__":
    main()
